import os
import pytest
from pathlib import Path
from pptx import Presentation

from backend.chats.presentation.schemas import PresentationPlan, SlideSpec, SlidePluginParagraph, SlidePluginImage, SlidePluginBullets, SlidePluginChart, SlidePluginTable
from backend.chats.presentation.geometry import SAFE_CONTENT_BOTTOM, calculate_available_content_height, MixedLayoutResolver, Box
from backend.chats.presentation.planner import analyze_content_density, split_overdense_slides, PARAGRAPH_MIN_FONT_SIZE, BULLET_MIN_FONT_SIZE
from backend.chats.presentation.renderers.ppt_renderer import PptRenderer


def test_safe_content_bottom_constants():
    assert SAFE_CONTENT_BOTTOM == 6.70
    assert calculate_available_content_height(1.4) == 5.30
    assert calculate_available_content_height(2.0) == 4.70


def create_test_plan(slides):
    return PresentationPlan(
        title="Layout Engine Test Suite",
        theme={"name": "modern_corporate", "background": "#FFFFFF", "text": "#1E293B", "accent": "#3B82F6"},
        slides=slides,
    )


def inspect_generated_pptx(prs: Presentation):
    """
    Asserts that NO content text box or shape in any slide has a bottom position exceeding SAFE_CONTENT_BOTTOM + 0.15
    and that text frames and images stay inside safe bounds (above footer zone starting at 6.85).
    """
    for idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame or shape.shape_type == 13: # Picture or text shape
                top_in = shape.top / 914400.0 # EMU to Inches
                height_in = shape.height / 914400.0
                bottom_in = round(top_in + height_in, 2)
                # Content shapes start below title (top_in >= 1.0) and above footer zone (top_in < 6.80)
                if 1.0 <= top_in < 6.80:
                    assert bottom_in <= SAFE_CONTENT_BOTTOM + 0.15, (
                        f"Slide {idx + 1} content shape '{getattr(shape, 'text', 'visual')[:30]}' "
                        f"top={top_in}in, height={height_in}in, bottom={bottom_in}in exceeds safe content bottom {SAFE_CONTENT_BOTTOM}in"
                    )


# CASE 1: Short paragraph + image
def test_case_1_short_paragraph_and_image(tmp_path):
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Short Paragraph & Image",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": "Short intro paragraph detailing core concept."}),
            SlidePluginImage(type="image", data={"caption": "Concept Visual"}),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 2: Medium paragraph + image
def test_case_2_medium_paragraph_and_image(tmp_path):
    text = (
        "As Diwali continues to expand its footprint across international capitals and corporate boardrooms, "
        "its core message of unity and resilience remains timeless. Future celebrations will increasingly prioritize "
        "sustainable practices and green energy illumination over resource-intensive consumption."
    )
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Strategic Takeaways & Future Horizon",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text}),
            SlidePluginImage(type="image", data={"caption": "Diwali Illumination"}),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 3: Very long paragraph + image
def test_case_3_very_long_paragraph_and_image(tmp_path):
    text = (
        "As Diwali continues to expand its footprint across international capitals and corporate boardrooms, "
        "its core message of unity and resilience remains timeless. Future celebrations will increasingly prioritize "
        "sustainable practices, green energy illumination, and philanthropic engagement over resource-intensive "
        "consumption. By fostering inclusive environments in both educational institutions and global enterprises, "
        "leaders can leverage the festival's universal values to drive cultural awareness, cross-cultural collaboration, "
        "and long-term community empowerment across global markets."
    )
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Strategic Takeaways & Future Horizon",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text}),
            SlidePluginImage(type="image", data={"caption": "Strategic Horizon"}),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 4: Long paragraph without image
def test_case_4_long_paragraph_without_image(tmp_path):
    text = (
        "Comprehensive enterprise architecture requires decoupled microservices, resilient fault tolerance, "
        "distributed caching layer, zero-trust security policies, automated CI/CD deployment pipelines, and continuous "
        "observability to achieve 99.999% uptime SLAs under peak seasonal traffic demands."
    )
    slide = SlideSpec(
        layout="title_content",
        title="Enterprise Architecture Guidelines",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text})
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 5: Multiple paragraphs + image
def test_case_5_multiple_paragraphs_and_image(tmp_path):
    text1 = "First paragraph providing high-level executive summary and domain context."
    text2 = "Second paragraph outlining implementation steps, governance controls, and performance metrics."
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Executive Summary & Implementation",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text1}),
            SlidePluginParagraph(type="paragraph", data={"text": text2}),
            SlidePluginImage(type="image", data={"caption": "Implementation Workflow"}),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 6: Paragraph + bullets + image
def test_case_6_paragraph_bullets_and_image(tmp_path):
    text = "Overview paragraph introducing key operational milestones."
    bullets = ["Milestone 1: Platform launch", "Milestone 2: User scaling", "Milestone 3: Enterprise integration"]
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Operational Roadmap",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text}),
            SlidePluginBullets(type="bullets", data={"points": bullets}),
            SlidePluginImage(type="image", data={"caption": "Roadmap Overview"}),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 7: Very long bullet list
def test_case_7_very_long_bullet_list(tmp_path):
    bullets = [f"Strategic Priority #{i}: Expand operational readiness and audit controls across regional business units." for i in range(1, 9)]
    slide = SlideSpec(
        layout="bullets_slide",
        title="Strategic Priorities Checklist",
        plugins=[
            SlidePluginBullets(type="bullets", data={"points": bullets})
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 8: Chart + paragraph
def test_case_8_chart_and_paragraph(tmp_path):
    text = "Quarterly revenue growth trajectory demonstrates a 34% YoY increase driven by subscription expansion."
    chart_payload = {
        "chart_type": "column",
        "title": "Revenue Growth",
        "categories": ["Q1", "Q2", "Q3", "Q4"],
        "values": [120, 150, 180, 220],
        "series_name": "Revenue ($M)"
    }
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Financial Trajectory",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text}),
            SlidePluginChart(type="chart", data=chart_payload),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 9: Table + paragraph
def test_case_9_table_and_paragraph(tmp_path):
    text = "Comparison matrix highlighting feature availability across pricing tiers."
    table_payload = {
        "headers": ["Feature", "Basic", "Pro", "Enterprise"],
        "rows": [
            ["Custom Branding", "No", "Yes", "Yes"],
            ["API Access", "No", "Limited", "Full"],
            ["SLA Uptime", "99.0%", "99.9%", "99.999%"],
        ]
    }
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Feature Tier Comparison",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": text}),
            SlidePluginTable(type="table", data=table_payload),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 10: Extremely large content triggering automatic slide splitting
def test_case_10_extreme_content_slide_splitting(tmp_path):
    huge_paragraph = " ".join([
        f"Sentence {i}: This is an extremely detailed sentence explaining complex domain architecture, operational protocols, compliance guidelines, and long-term risk mitigation strategies."
        for i in range(1, 15)
    ])
    slide = SlideSpec(
        layout="mixed_content_slide",
        title="Massive Architectural Specification",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": huge_paragraph}),
            SlidePluginImage(type="image", data={"caption": "Architecture Diagram"}),
        ]
    )
    plan = create_test_plan([slide])
    split_plan = split_overdense_slides(plan)
    assert len(split_plan.slides) >= 2, "Overdense slide should automatically split into 2 slides"
    
    renderer = PptRenderer()
    prs = renderer.render(plan) # PptRenderer.render automatically splits as well
    assert len(prs.slides) >= 2
    inspect_generated_pptx(prs)


# CASE 11: Very short content where excessive empty space is avoided
def test_case_11_very_short_content(tmp_path):
    slide = SlideSpec(
        layout="title_content",
        title="Core Vision",
        plugins=[
            SlidePluginParagraph(type="paragraph", data={"text": "Empowering teams through intelligent presentation automation."})
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)


# CASE 12: Long image caption handling
def test_case_12_long_image_caption(tmp_path):
    long_caption = "fig:- Comprehensive Global Distribution Infrastructure Map Illustrating Multi-Region Edge Caching Nodes and Real-Time Event Telemetry Stream Synchronization Protocols"
    slide = SlideSpec(
        layout="image_slide",
        title="Global Infrastructure Map",
        plugins=[
            SlidePluginImage(type="image", data={"caption": long_caption, "title": "Infrastructure Map"}),
        ]
    )
    plan = create_test_plan([slide])
    renderer = PptRenderer()
    prs = renderer.render(plan)
    inspect_generated_pptx(prs)
