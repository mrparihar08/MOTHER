import os
import tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from backend.presentation.scripts.generate_master_template import (
    GEOMETRY,
    GeometryConfig,
    ThemeConfig,
    LayoutResolver,
    PPTBuilder,
    create_master_template,
    calculate_luminance,
    calculate_contrast_ratio,
    validate_text_contrast,
    validate_shape_bounds,
    add_blank_slide,
)


def test_geometry_and_safe_bounds():
    assert GEOMETRY.SLIDE_WIDTH == 13.333
    assert GEOMETRY.SLIDE_HEIGHT == 7.5
    assert GEOMETRY.MARGIN_X == 0.7
    assert GEOMETRY.MARGIN_Y == 0.5

    # Safe bounds validation
    assert validate_shape_bounds(0.7, 1.4, 11.9, 5.2, "Test") is True
    assert validate_shape_bounds(-0.1, 1.4, 11.9, 5.2, "Negative x") is False
    assert validate_shape_bounds(0.7, -0.5, 11.9, 5.2, "Negative y") is False
    assert validate_shape_bounds(0.7, 1.4, 0, 5.2, "Zero width") is False
    assert validate_shape_bounds(0.7, 1.4, 14.0, 5.2, "Overflow width") is False

    # Multi-column grid distribution
    cols2 = GEOMETRY.get_column_rects(2, gap=0.5)
    assert len(cols2) == 2
    for c in cols2:
        assert validate_shape_bounds(*c) is True

    cols4 = GEOMETRY.get_column_rects(4, gap=0.3)
    assert len(cols4) == 4
    for c in cols4:
        assert validate_shape_bounds(*c) is True


def test_luminance_and_contrast_validation():
    dark_slate = RGBColor(15, 23, 42)
    white = RGBColor(255, 255, 255)
    light_slate = RGBColor(248, 250, 252)

    assert calculate_luminance(white) > 250
    assert calculate_luminance(dark_slate) < 30

    ratio = calculate_contrast_ratio(white, dark_slate)
    assert ratio > 4.5  # Legible high contrast
    assert ratio >= 7.0  # WCAG AAA level contrast

    assert validate_text_contrast(white, dark_slate, min_ratio=4.5) is True
    assert validate_text_contrast(light_slate, white, min_ratio=4.5) is False


def test_layout_resolver():
    assert LayoutResolver.resolve_layout({"is_cover": True}) == "cover"
    assert LayoutResolver.resolve_layout({"is_section": True}) == "section"
    assert LayoutResolver.resolve_layout({"milestones": [{"title": "M1"}]}) == "timeline"
    assert LayoutResolver.resolve_layout({"steps": [{"title": "S1"}]}) == "process"
    assert LayoutResolver.resolve_layout({"quote": "A quote"}) == "quote"
    assert LayoutResolver.resolve_layout({"metrics": [{"value": "100"}]}) == "statistics"
    assert LayoutResolver.resolve_layout({"headers": ["A"], "rows": [["1"]]}) == "table"
    assert LayoutResolver.resolve_layout({"chart_data": [1, 2, 3]}) == "chart"
    assert LayoutResolver.resolve_layout({"left_column": {}, "right_column": {}}) == "two_column"
    assert LayoutResolver.resolve_layout({"cards": [1, 2, 3]}) == "three_column"
    assert LayoutResolver.resolve_layout({"image": "img.jpg", "paragraph": "Text", "bullets": ["b1"]}) == "mixed_content"
    assert LayoutResolver.resolve_layout({"image": "img.jpg", "paragraph": "Text"}) == "image_text"
    assert LayoutResolver.resolve_layout({"paragraph": "Text", "bullets": ["b1"]}) == "mixed_content"
    assert LayoutResolver.resolve_layout({"paragraph": "Only text"}) == "title_content"


def test_create_master_template_file_and_slides():
    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = os.path.join(tmp_dir, "base_template.pptx")
        res_path = create_master_template(out_path)
        assert os.path.isfile(res_path)
        assert res_path == str(Path(out_path).resolve())

        # Load and verify PowerPoint structure
        prs = Presentation(res_path)
        assert prs.slide_width == Inches(13.333)
        assert prs.slide_height == Inches(7.5)
        assert len(prs.slides) == 14  # All 14 layout types present

        for slide in prs.slides:
            assert len(slide.shapes) > 0


def test_pptbuilder_individual_layouts_and_data_rendering():
    builder = PPTBuilder()
    prs = builder.create_presentation()

    # Test all individual layout renderers
    builder.add_cover(prs, title="Cover Test")
    builder.add_title_content(prs, title="Content Test")
    builder.add_two_column(prs, title="Two Column Test")
    builder.add_three_column(prs, title="Three Column Test")
    builder.add_section(prs, section_title="Section Test")
    builder.add_quote(prs, quote_text="Test Quote")
    builder.add_statistics(prs, title="Stats Test")
    builder.add_comparison(prs, title="Comparison Test")
    builder.add_timeline(prs, title="Timeline Test")
    builder.add_process(prs, title="Process Test")
    builder.add_image_text(prs, title="Image Text Test")
    builder.add_chart(prs, title="Chart Test")
    builder.add_table(prs, title="Table Test")
    builder.add_mixed_content(prs, title="Mixed Test")

    assert len(prs.slides) == 14

    # Test dynamic structured data rendering
    dynamic_slide_data = {
        "title": "Dynamic AI Slide",
        "subtitle": "Generated via LLM",
        "steps": [
            {"step": "1", "title": "Prompt", "desc": "User enters idea"},
            {"step": "2", "title": "Planner", "desc": "Semantic planner chooses layout"},
            {"step": "3", "title": "Renderer", "desc": "PPTBuilder renders shapes"},
        ],
    }
    builder.render_slide_from_data(prs, dynamic_slide_data)
    assert len(prs.slides) == 15
