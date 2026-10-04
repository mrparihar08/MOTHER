import os
import tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from backend.chats.presentation.scripts.templates import (
    TEMPLATE_ARCHETYPES,
    TEMPLATE_PRESETS,
    ThemeConfig,
    GeometrySystem,
    TypographySystem,
    SlideRenderer,
    build_template_pptx,
    get_template_preset,
    list_available_templates,
    to_hex_str,
    to_rgb,
    is_light_rgb,
    validate_layout,
)


def test_theme_config_creation_and_luminance():
    # Light theme test
    light_theme = ThemeConfig(
        id="test_light",
        name="Test Light Theme",
        bg=RGBColor(255, 255, 255),
        card_bg=RGBColor(245, 245, 245),
        text_dark=RGBColor(15, 23, 42),
        text_light=RGBColor(255, 255, 255),
    )
    assert light_theme.is_bg_light() is True
    assert light_theme.is_card_bg_light() is True
    assert light_theme.get_text_for_bg(light_theme.bg) == light_theme.text_dark

    # Dark theme test
    dark_theme = ThemeConfig(
        id="test_dark",
        name="Test Dark Theme",
        bg=RGBColor(15, 23, 42),
        card_bg=RGBColor(30, 41, 59),
        text_dark=RGBColor(15, 23, 42),
        text_light=RGBColor(255, 255, 255),
    )
    assert dark_theme.is_bg_light() is False
    assert dark_theme.is_card_bg_light() is False
    assert dark_theme.get_text_for_bg(dark_theme.bg) == dark_theme.text_light


def test_geometry_system_16_9():
    assert GeometrySystem.SLIDE_WIDTH_INCHES == 13.333
    assert GeometrySystem.SLIDE_HEIGHT_INCHES == 7.5

    rect = GeometrySystem.get_content_rect(style_type="modern")
    assert rect[0] == 0.8
    assert rect[1] == 1.3
    assert abs(rect[2] - (13.333 - 1.6)) < 1e-4
    assert rect[3] == 5.6
    assert GeometrySystem.validate_bounds(*rect) is True

    # Sidebar layout rect
    sidebar_rect = GeometrySystem.get_content_rect(style_type="sidebar", sidebar_w=2.5)
    assert sidebar_rect[0] == 2.9
    assert GeometrySystem.validate_bounds(*sidebar_rect) is True

    # 3-column split rects
    cols = GeometrySystem.get_column_rects(3, y=1.5, h=5.0, gap=0.4)
    assert len(cols) == 3
    for col in cols:
        assert GeometrySystem.validate_bounds(*col) is True


def test_typography_system():
    theme = ThemeConfig(id="test", name="Test")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    p = tbox.text_frame.paragraphs[0]

    TypographySystem.apply_title(p, "Main Title", theme)
    assert p.text == "Main Title"
    assert p.font.bold is True
    assert p.font.size == TypographySystem.SLIDE_TITLE_SIZE

    TypographySystem.apply_body(p, "Body description", theme)
    assert p.text == "Body description"
    assert p.font.size == TypographySystem.BODY_SIZE


def test_get_template_preset_and_listing():
    # Check that all 16 archetypes are present in TEMPLATE_ARCHETYPES
    archetype_keys = [
        "ion_boardroom", "berlin_executive", "quotable_teal", "geometric_block",
        "urban_monochrome", "crop_frame", "circuit_tech", "celestial_night",
        "artistic_neon", "atlas_bold", "organic_pastel", "dividend_burgundy",
        "savon_classic", "wood_type", "sidebar_executive", "modern_glassmorphism"
    ]
    for key in archetype_keys:
        assert key in TEMPLATE_ARCHETYPES
        assert key in TEMPLATE_PRESETS
        preset = get_template_preset(key)
        assert preset is not None
        assert "name" in preset
        assert "style_type" in preset

    # Test list_available_templates
    templates = list_available_templates()
    assert len(templates) >= 22
    for item in templates:
        assert "id" in item
        assert "name" in item
        assert "bg" in item
        assert item["bg"].startswith("#")
        assert "font_color" in item


def test_build_all_16_template_archetypes_pptx():
    with tempfile.TemporaryDirectory() as tmp_dir:
        for key, cfg in TEMPLATE_ARCHETYPES.items():
            pptx_path = build_template_pptx(key, cfg, output_dir=tmp_dir)
            assert os.path.isfile(pptx_path)

            # Open presentation and verify structure
            prs = Presentation(pptx_path)
            assert prs.slide_width == Inches(13.333)
            assert prs.slide_height == Inches(7.5)
            assert len(prs.slides) == 9  # 9 slide layouts created

            # Verify slides have no default placeholder collisions
            for slide in prs.slides:
                # Should have custom shapes rendered
                assert len(slide.shapes) > 0


def test_build_all_presets_pptx():
    with tempfile.TemporaryDirectory() as tmp_dir:
        presets_to_test = ["base_template", "corporate_light", "modern_dark", "emerald_nature", "executive_gold", "cyber_neon"]
        for key in presets_to_test:
            cfg = TEMPLATE_PRESETS[key]
            pptx_path = build_template_pptx(key, cfg, output_dir=tmp_dir)
            assert os.path.isfile(pptx_path)

            prs = Presentation(pptx_path)
            assert prs.slide_width == Inches(13.333)
            assert prs.slide_height == Inches(7.5)
            assert len(prs.slides) == 9
