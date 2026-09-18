from __future__ import annotations

import logging
import math
import os
import re
import requests
from pathlib import Path
from typing import Any, Dict, List, Optional
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from backend.chats.services.unsplash_service import fetch_unsplash_image, fetch_unsplash_url
from backend.chats.presentation.schemas import PresentationPlan, SlideSpec
from backend.chats.presentation.themes import (
    get_theme_palette,
    apply_background_theme,
    hex_to_rgb,
    is_light_color,
    ensure_readable_text_color,
)
from backend.chats.presentation.geometry import (
    Box,
    as_box,
    get_layout_registry,
    LayoutSpec,
)
from backend.chats.presentation.planner import (
    normalize_whitespace,
    clean_ai_instructions,
    safe_filename,
    best_font_size_for_bullets,
    best_font_size_for_paragraph,
    parse_number,
    safe_list,
    detect_bullet_style,
    format_bullet_prefix,
    detect_diagram_type,
    build_dynamic_diagram_badge,
)

logger = logging.getLogger(__name__)

DEFAULT_TEMPLATE_FILE = os.getenv("PPT_TEMPLATE_FILE", "./templates/base_template.pptx")
ASSET_DIR = Path(os.getenv("PPT_ASSET_DIR", "./assets")).resolve()
ALLOW_ABSOLUTE_IMAGE_PATHS = os.getenv("PPT_ALLOW_ABSOLUTE_IMAGE_PATHS", "false").lower() == "true"
ASSET_DIR.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------
# Rendering Helper Functions
# ---------------------------------------------------------------------

def ensure_template_prs(template_file: str) -> Presentation:
    path = Path(template_file)
    prs = None
    if path.exists():
        try:
            prs = Presentation(str(path))
        except Exception as exc:
            logger.warning("Failed to load template %s: %s", path, exc)
    if prs is None:
        prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def set_run_style(run, font_size: int, bold: bool = False, color: Optional[RGBColor] = None) -> None:
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_text_frame(tf, *, font_size: int, color: Optional[RGBColor] = None, bold: bool = False) -> None:
    try:
        tf.word_wrap = True
    except Exception:
        pass
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    try:
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.12)
        tf.margin_bottom = Inches(0.12)
    except Exception:
        pass

    for p in tf.paragraphs:
        for run in p.runs:
            set_run_style(run, font_size=font_size, bold=bold, color=color)


def find_placeholder_by_types(slide, placeholder_types: tuple) -> Optional[Any]:
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type in placeholder_types:
                return shape
        except Exception:
            continue
    return None


def set_shape_text(shape, text: str, font_size: int = 20, bold: bool = False, color: Optional[RGBColor] = None) -> None:
    if not hasattr(shape, "text_frame"):
        return
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_style(run, font_size=font_size, bold=bold, color=color)
    try:
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass


def add_textbox(slide, left, top, width, height, text: str, font_size: int = 20, bold: bool = False, color: Optional[RGBColor] = None) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_style(run, font_size=font_size, bold=bold, color=color)
    try:
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass


def write_text_or_fallback(
    slide,
    placeholder_types: tuple,
    text: str,
    *,
    fallback_left: float,
    fallback_top: float,
    fallback_width: float,
    fallback_height: float,
    font_size: int,
    bold: bool = False,
    color: Optional[RGBColor] = None,
) -> None:
    shape = find_placeholder_by_types(slide, placeholder_types)
    if shape is not None:
        try:
            set_shape_text(shape, text, font_size=font_size, bold=bold, color=color)
            return
        except Exception:
            pass

    add_textbox(
        slide,
        Inches(fallback_left),
        Inches(fallback_top),
        Inches(fallback_width),
        Inches(fallback_height),
        text,
        font_size=font_size,
        bold=bold,
        color=color,
    )


def set_slide_notes(slide, notes: str) -> None:
    notes = normalize_whitespace(notes)
    if not notes:
        return
    try:
        ns = slide.notes_slide
        tf = ns.notes_text_frame
        tf.clear()
        tf.text = notes
    except Exception:
        pass


from backend.chats.presentation.services.security import is_safe_url, is_safe_image_path

def sanitize_image_path(path_text: str) -> Optional[str]:
    path_text = normalize_whitespace(path_text)
    if not path_text:
        return None

    if path_text.startswith("http://") or path_text.startswith("https://"):
        if not is_safe_url(path_text):
            logger.warning("Blocked unsafe URL: %s", path_text)
            return None

        url_clean = path_text.split("?")[0].rstrip("/")
        file_part = url_clean.split("/")[-1] or "web_image"
        safe_name = safe_filename(file_part)
        target_file = ASSET_DIR / f"download_{safe_name}.jpg"
        if target_file.exists() and target_file.stat().st_size > 1000:
            return str(target_file)
        try:
            resp = requests.get(path_text, timeout=10, stream=True)
            if resp.status_code == 200:
                content_len = int(resp.headers.get("Content-Length", 0))
                if content_len > 15 * 1024 * 1024:
                    logger.warning("Downloaded image exceeded max size limit: %s", path_text)
                    return None
                data = resp.content
                if len(data) > 1000:
                    target_file.write_bytes(data)
                    return str(target_file)
        except Exception as exc:
            logger.warning("Failed to download image from URL %s: %s", path_text, exc)

    safe_local = is_safe_image_path(path_text)
    if safe_local:
        return str(safe_local)
    return None


def fetch_unsplash_image_for_topic(topic_query: str) -> Optional[str]:
    return fetch_unsplash_image(topic_query)


def add_card_container(slide, box: Box, palette: Dict[str, RGBColor], border_color: Optional[RGBColor] = None) -> Any:
    """Adds a modern rounded card container shape behind content blocks for visual structure."""
    try:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(box.left),
            Inches(box.top),
            Inches(box.width),
            Inches(box.height)
        )
        card.fill.solid()
        bg_rgb = palette.get("card_bg") or palette.get("table_row_bg1") or palette["background"]
        card.fill.fore_color.rgb = bg_rgb

        b_color = border_color or palette.get("card_border") or palette.get("accent")
        if b_color:
            card.line.color.rgb = b_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card
    except Exception as exc:
        logger.warning("Failed to render card container: %s", exc)
        return None


# ---------------------------------------------------------------------
# Base & Derived Plugin Implementations
# ---------------------------------------------------------------------

class BasePlugin:
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        pass

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        self.apply(slide, plan, theme_name=None)
        return current_y


class TextPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        text = normalize_whitespace(plan.get("text", "") or plan.get("subtitle", "") or plan.get("title", ""))
        if not text:
            return
        box_spec = as_box(plan, Box(0.8, 1.5, 11.7, 0.5))
        box = slide.shapes.add_textbox(Inches(box_spec.left), Inches(box_spec.top), Inches(box_spec.width), Inches(box_spec.height))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        user_font = plan.get("font_size")
        default_size = 23 if plan.get("type") == "subtitle" else 29
        font_size = int(user_font) if user_font and str(user_font).isdigit() else default_size
        p.font.size = Pt(font_size)
        p.font.bold = True
        custom_color = plan.get("font_color") or plan.get("color")
        p.font.color.rgb = hex_to_rgb(custom_color) if custom_color else palette["accent"]

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        text = normalize_whitespace(plan.get("text", "") or plan.get("subtitle", "") or plan.get("title", ""))
        if not text:
            return current_y
        box = slide.shapes.add_textbox(Inches(left_margin), Inches(current_y), Inches(content_width), Inches(0.5))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        user_font = plan.get("font_size")
        default_size = 23 if plan.get("type") == "subtitle" else 29
        font_size = int(user_font) if user_font and str(user_font).isdigit() else default_size
        p.font.size = Pt(font_size)
        p.font.bold = True
        custom_color = plan.get("font_color") or plan.get("color")
        p.font.color.rgb = hex_to_rgb(custom_color) if custom_color else palette["accent"]
        return current_y + 0.65


class ParagraphPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        text = normalize_whitespace(plan.get("text", ""))
        if not text:
            return

        user_font = plan.get("font_size")
        font_size = int(user_font) if user_font and str(user_font).isdigit() else best_font_size_for_paragraph(text, base=14)

        default_height = 0.6 if len(text) < 120 else (0.8 if len(text) < 250 else 1.1)
        box_spec = as_box(plan, Box(0.8, 1.5, 5.6, default_height))

        box = slide.shapes.add_textbox(Inches(box_spec.left), Inches(box_spec.top), Inches(box_spec.width), Inches(box_spec.height))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True

        pts = safe_list(plan.get("points"))
        if not pts and text and ("\n" in text or any(text.startswith(p) for p in ("•", "-", "*", "1.", "2.", "✓", "➔"))):
            lines = [l.strip() for l in text.split("\n") if l.strip()]
            if len(lines) > 1 or (lines and any(lines[0].startswith(p) for p in ("•", "-", "*", "1.", "2.", "✓", "➔"))):
                pts = [re.sub(r"^[•\-*✓➔\d+\.\s]+", "", l).strip() for l in lines]

        custom_color = plan.get("font_color") or plan.get("color")
        text_color = hex_to_rgb(custom_color) if custom_color else palette["text"]

        if pts:
            bullet_font = int(user_font) if user_font and str(user_font).isdigit() else best_font_size_for_bullets(pts, base=14)
            b_style = plan.get("bullet_style") or plan.get("list_style") or "disc"
            for idx, pt in enumerate(pts):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                prefix = format_bullet_prefix(b_style, idx, pts)
                p.text = f"{prefix} {pt}"
                p.space_after = Pt(2)
            configure_text_frame(tf, font_size=bullet_font, color=text_color)
        else:
            tf.text = text
            configure_text_frame(tf, font_size=font_size, color=text_color)

        alignment = str(plan.get("alignment", plan.get("align", "left"))).lower()
        align_map = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY, "left": PP_ALIGN.LEFT}
        v_align_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        if alignment in align_map:
            for p in tf.paragraphs:
                p.alignment = align_map[alignment]
        if alignment in v_align_map:
            tf.vertical_anchor = v_align_map[alignment]

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        text = normalize_whitespace(plan.get("text", ""))
        if not text:
            return current_y

        user_font = plan.get("font_size")
        font_size = int(user_font) if user_font and str(user_font).isdigit() else best_font_size_for_paragraph(text, base=14)

        default_height = 0.6 if len(text) < 120 else (0.8 if len(text) < 250 else 1.1)
        box_spec = as_box(plan, Box(left_margin, current_y, content_width, default_height))

        box = slide.shapes.add_textbox(Inches(box_spec.left), Inches(box_spec.top), Inches(box_spec.width), Inches(box_spec.height))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.text = text
        custom_color = plan.get("font_color") or plan.get("color")
        text_color = hex_to_rgb(custom_color) if custom_color else palette["text"]
        configure_text_frame(tf, font_size=font_size, color=text_color)

        alignment = str(plan.get("alignment", plan.get("align", "left"))).lower()
        align_map = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY, "left": PP_ALIGN.LEFT}
        v_align_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        if alignment in align_map:
            for p in tf.paragraphs:
                p.alignment = align_map[alignment]
        if alignment in v_align_map:
            tf.vertical_anchor = v_align_map[alignment]

        return max(current_y, box_spec.top) + box_spec.height + 0.15


class BulletsPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        points = safe_list(plan.get("points"))
        if not points:
            return

        user_font = plan.get("font_size")
        bullet_font = int(user_font) if user_font and str(user_font).isdigit() else best_font_size_for_bullets(points, base=14)

        default_height = max(0.6, 0.28 * len(points))
        top_pos = float(plan.get("top", 1.8))
        box_spec = as_box(plan, Box(0.9, top_pos, 8.5, default_height))

        if plan.get("show_card", False):
            add_card_container(slide, box_spec, palette)

        box = slide.shapes.add_textbox(Inches(box_spec.left), Inches(box_spec.top), Inches(box_spec.width), Inches(box_spec.height))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        bullet_style = plan.get("bullet_style") or plan.get("list_style") or "auto"
        for idx, point in enumerate(points):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            prefix = format_bullet_prefix(bullet_style, idx, points)
            p.text = f"{prefix} {point}"
            p.level = 0
            p.space_after = Pt(2)

        custom_color = plan.get("font_color") or plan.get("color")
        text_color = hex_to_rgb(custom_color) if custom_color else palette["text"]
        configure_text_frame(tf, font_size=bullet_font, color=text_color)

        alignment = str(plan.get("alignment", plan.get("align", "left"))).lower()
        align_map = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY, "left": PP_ALIGN.LEFT}
        v_align_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        if alignment in align_map:
            for p in tf.paragraphs:
                p.alignment = align_map[alignment]
        if alignment in v_align_map:
            tf.vertical_anchor = v_align_map[alignment]

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        points = safe_list(plan.get("points"))
        if not points:
            return current_y

        user_font = plan.get("font_size")
        bullet_font = int(user_font) if user_font and str(user_font).isdigit() else best_font_size_for_bullets(points, base=14)

        default_height = max(0.6, 0.28 * len(points))
        box_spec = as_box(plan, Box(left_margin, current_y, content_width, default_height))

        if plan.get("show_card", False):
            add_card_container(slide, box_spec, palette)

        box = slide.shapes.add_textbox(Inches(box_spec.left), Inches(box_spec.top), Inches(box_spec.width), Inches(box_spec.height))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        bullet_style = plan.get("bullet_style") or plan.get("list_style") or "auto"
        for idx, point in enumerate(points):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            prefix = format_bullet_prefix(bullet_style, idx, points)
            p.text = f"{prefix} {point}"
            p.level = 0
            p.space_after = Pt(2)

        custom_color = plan.get("font_color") or plan.get("color")
        text_color = hex_to_rgb(custom_color) if custom_color else palette["text"]
        configure_text_frame(tf, font_size=bullet_font, color=text_color)

        alignment = str(plan.get("alignment", plan.get("align", "left"))).lower()
        align_map = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY, "left": PP_ALIGN.LEFT}
        v_align_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        if alignment in align_map:
            for p in tf.paragraphs:
                p.alignment = align_map[alignment]
        if alignment in v_align_map:
            tf.vertical_anchor = v_align_map[alignment]

        return max(current_y, box_spec.top) + box_spec.height + 0.20


class ChartPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        chart_type = str(plan.get("chart_type", "column")).lower()
        categories = list(plan.get("categories", []))
        values = list(plan.get("values", []))
        series_map = plan.get("series_map") or {}
        series_name = normalize_whitespace(plan.get("series_name", "Usage"))
        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.9, top_pos, 8.5, 3.8))
        safe_top = min(raw_box.top, 4.5)
        safe_height = min(raw_box.height, round(6.8 - safe_top, 2))
        box = Box(raw_box.left, safe_top, raw_box.width, max(1.5, safe_height))

        chart_data = CategoryChartData()
        if series_map:
            has_nonzero = False
            for mapping in series_map.values():
                if isinstance(mapping, dict):
                    for v in mapping.values():
                        num = parse_number(re.sub(r"[^\d.-]", "", str(v))) if not isinstance(v, (int, float)) else float(v)
                        if num and num != 0.0:
                            has_nonzero = True
                            break
                elif isinstance(mapping, (list, tuple)):
                    for v in mapping:
                        num = parse_number(re.sub(r"[^\d.-]", "", str(v))) if not isinstance(v, (int, float)) else float(v)
                        if num and num != 0.0:
                            has_nonzero = True
                            break
            if not has_nonzero:
                series_map = None

        if series_map:
            if not categories:
                seen: List[str] = []
                for mapping in series_map.values():
                    if isinstance(mapping, dict):
                        for cat in mapping.keys():
                            if cat not in seen:
                                seen.append(cat)
                categories = seen
            if not categories:
                categories = ["Phase 1", "Phase 2", "Phase 3", "Phase 4"]
            chart_data.categories = categories
            for s_name, mapping in series_map.items():
                if isinstance(mapping, dict):
                    series_values = []
                    for cat in categories:
                        val = mapping.get(cat, 0)
                        num = parse_number(re.sub(r"[^\d.-]", "", str(val))) if not isinstance(val, (int, float)) else float(val)
                        series_values.append(num if num is not None else 0.0)
                    chart_data.add_series(str(s_name), series_values)
                elif isinstance(mapping, (list, tuple)):
                    series_values = []
                    for v in mapping[:len(categories)]:
                        num = parse_number(re.sub(r"[^\d.-]", "", str(v))) if not isinstance(val, (int, float)) else float(val)
                        series_values.append(num if num is not None else 0.0)
                    if len(series_values) < len(categories):
                        series_values += [0.0] * (len(categories) - len(series_values))
                    chart_data.add_series(str(s_name), series_values)
        else:
            clean_values = []
            for v in values:
                num = parse_number(re.sub(r"[^\d.-]", "", str(v))) if not isinstance(v, (int, float)) else float(v)
                clean_values.append(num if num is not None else 0.0)

            if not clean_values or all(v == 0.0 for v in clean_values):
                categories = categories if categories else ["Phase 1 (Baseline)", "Phase 2 (Growth)", "Phase 3 (Scale)", "Phase 4 (Maturity)"]
                values = [round(28.5 * (i + 1) * 1.15, 1) for i in range(len(categories))]
                if "[Illustrative Data]" not in series_name:
                    series_name = f"{series_name} [Illustrative Data]"
            else:
                values = clean_values
                if not categories:
                    categories = [f"Phase {i+1}" for i in range(len(values))]

            n = min(len(categories), len(values))
            categories = categories[:n]
            values = values[:n]

            chart_data.categories = categories
            chart_data.add_series(series_name, values)

        chart_kind = XL_CHART_TYPE.COLUMN_CLUSTERED
        if chart_type in {"line", "trend"}:
            chart_kind = XL_CHART_TYPE.LINE_MARKERS
        elif chart_type in {"bar", "bar_horizontal"}:
            chart_kind = XL_CHART_TYPE.BAR_CLUSTERED
        elif chart_type in {"pie"}:
            chart_kind = XL_CHART_TYPE.PIE
        elif chart_type in {"area"}:
            chart_kind = XL_CHART_TYPE.AREA
        elif chart_type in {"donut", "doughnut"}:
            chart_kind = XL_CHART_TYPE.DOUGHNUT
        elif chart_type in {"column"}:
            chart_kind = XL_CHART_TYPE.COLUMN_CLUSTERED
        elif chart_type in {"radar", "spider"}:
            chart_kind = getattr(XL_CHART_TYPE, "RADAR_FILLED", XL_CHART_TYPE.COLUMN_CLUSTERED)
        elif chart_type in {"gauge"}:
            chart_kind = XL_CHART_TYPE.DOUGHNUT
        elif chart_type in {"waterfall"}:
            chart_kind = getattr(XL_CHART_TYPE, "COLUMN_STACKED", XL_CHART_TYPE.COLUMN_CLUSTERED)

        try:
            chart_shape = slide.shapes.add_chart(chart_kind, Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height), chart_data)
            chart = chart_shape.chart

            title_text = plan.get("title")
            slide_title = plan.get("slide_title")
            if title_text and (title_text == slide_title or plan.get("show_title") is False):
                chart.has_title = False
            elif title_text and plan.get("show_title", True):
                chart.has_title = True
                chart.chart_title.text_frame.text = str(title_text)
                try:
                    p = chart.chart_title.text_frame.paragraphs[0]
                    p.font.color.rgb = palette["text"]
                    p.font.size = Pt(13)
                    p.font.bold = True
                except Exception:
                    pass

            show_legend = plan.get("show_legend", True)
            chart.has_legend = show_legend
            if show_legend:
                if "legend_position" in plan:
                    pos_key = str(plan["legend_position"]).lower()
                    pos_map = {
                        "top": XL_LEGEND_POSITION.TOP,
                        "bottom": XL_LEGEND_POSITION.BOTTOM,
                        "left": XL_LEGEND_POSITION.LEFT,
                        "right": XL_LEGEND_POSITION.RIGHT,
                    }
                    if pos_key in pos_map:
                        chart.legend.position = pos_map[pos_key]
                try:
                    chart.legend.font.color.rgb = palette["text"]
                    chart.legend.font.size = Pt(10)
                except Exception:
                    pass

            try:
                if hasattr(chart, "category_axis"):
                    chart.category_axis.tick_labels.font.color.rgb = palette["text"]
                    chart.category_axis.tick_labels.font.size = Pt(10)
            except Exception:
                pass

            try:
                if hasattr(chart, "value_axis"):
                    chart.value_axis.tick_labels.font.color.rgb = palette["text"]
                    chart.value_axis.tick_labels.font.size = Pt(10)
            except Exception:
                pass

            show_data_labels = plan.get("show_data_labels", False)
            if show_data_labels and len(chart.plots) > 0:
                plot = chart.plots[0]
                plot.has_data_labels = True
                try:
                    plot.data_labels.font.color.rgb = palette["text"]
                    plot.data_labels.font.size = Pt(10)
                except Exception:
                    pass

        except Exception as exc:
            logger.warning("Failed to render PPT chart: %s", exc)
            fallback_box = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(1.5))
            tf = fallback_box.text_frame
            tf.text = f"[{series_name} Chart]\nData: " + ", ".join(f"{c}: {v}" for c, v in zip(categories, values if not series_map else []))
            try:
                tf.paragraphs[0].runs[0].font.color.rgb = palette["text"]
            except Exception:
                pass

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        chart_w = min(11.0, content_width)
        chart_left = (13.333 - chart_w) / 2.0
        theme = plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": chart_left, "top": current_y, "width": chart_w, "height": 3.8}}, theme_name=theme)
        return current_y + 4.05


class DiagramPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        active_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(active_theme)
        diagram_text = plan.get("diagram", "") or plan.get("text", "") or "Input ➔ Process ➔ Output"
        selected_type = plan.get("diagram_type", "auto")
        context_text = plan.get("slide_title", "") or plan.get("title", "") or plan.get("caption", "") or ""
        diag_type = detect_diagram_type(diagram_text, selected_type, context_text)
        top_pos = float(plan.get("top", 1.5))
        box = as_box(plan, Box(0.8, top_pos, 11.7, 1.5))

        slide_title = str(plan.get("slide_title") or plan.get("title") or "").strip()
        custom_diag_title = str(plan.get("diagram_title") or plan.get("header") or plan.get("title") or "").strip()
        if custom_diag_title and custom_diag_title.lower() != slide_title.lower():
            dynamic_header = custom_diag_title.upper()
        else:
            dynamic_header = build_dynamic_diagram_badge(diag_type, context_text=context_text, slide_title=slide_title)

        stack_top = box.top
        if dynamic_header:
            badge_box = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(0.35))
            tf_b = badge_box.text_frame
            tf_b.word_wrap = True
            p_b = tf_b.paragraphs[0]
            p_b.alignment = PP_ALIGN.CENTER
            p_b.text = dynamic_header
            p_b.font.size = Pt(11)
            p_b.font.bold = True
            p_b.font.color.rgb = palette["accent"]
            stack_top = box.top + 0.35

        raw_steps = re.split(r"\s*(?:➔|➜|->|-->|→|⇒|\||\n|;)\s*", diagram_text)
        steps = []
        for s in raw_steps:
            cleaned = clean_ai_instructions(s).strip("[]()•-* ").strip()
            if "] [" in cleaned:
                for sub in cleaned.split("] ["):
                    sub_c = sub.strip("[]()•-* ").strip()
                    if sub_c:
                        steps.append(sub_c)
            elif cleaned:
                steps.append(cleaned)

        if len(steps) >= 2 and len(steps) <= 6:
            if diag_type == "architecture":
                layer_height = min(0.5, 1.2 / len(steps))
                gap = 0.08
                for i, step in enumerate(steps):
                    c_top = stack_top + i * (layer_height + gap)
                    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(box.left + 1.0), Inches(c_top), Inches(box.width - 2.0), Inches(layer_height))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = palette["accent"]
                    shape.line.color.rgb = palette["text"]
                    shape.line.width = Pt(1)
                    tf = shape.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = f"Layer {i + 1}: {step}"
                    set_run_style(run, font_size=11, bold=True, color=palette["background"])

            elif diag_type == "funnel":
                funnel_top = box.top + 0.4
                layer_height = min(0.45, 1.2 / len(steps))
                gap = 0.08
                total_w = box.width - 1.0
                for i, step in enumerate(steps):
                    w_factor = max(0.35, 1.0 - (i * (0.6 / max(1, len(steps) - 1))))
                    c_width = total_w * w_factor
                    c_left = box.left + (box.width - c_width) / 2.0
                    c_top = funnel_top + i * (layer_height + gap)
                    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(c_top), Inches(c_width), Inches(layer_height))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = palette["accent"]
                    shape.line.color.rgb = palette["text"]
                    shape.line.width = Pt(1)
                    tf = shape.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = f"Stage {i + 1}: {step}"
                    set_run_style(run, font_size=10, bold=True, color=palette["background"])

            elif diag_type == "pyramid":
                pyramid_top = box.top + 0.4
                layer_height = min(0.45, 1.2 / len(steps))
                gap = 0.08
                total_w = box.width - 1.0
                num_s = len(steps)
                for i, step in enumerate(reversed(steps)):
                    w_factor = max(0.35, 1.0 - (i * (0.6 / max(1, num_s - 1))))
                    c_width = total_w * w_factor
                    c_left = box.left + (box.width - c_width) / 2.0
                    c_top = pyramid_top + (num_s - 1 - i) * (layer_height + gap)
                    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(c_top), Inches(c_width), Inches(layer_height))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = palette["accent"]
                    shape.line.color.rgb = palette["text"]
                    shape.line.width = Pt(1)
                    tf = shape.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = f"Tier {num_s - i}: {step}"
                    set_run_style(run, font_size=10, bold=True, color=palette["background"])

            elif diag_type == "quadrant":
                q_top = box.top + 0.4
                q_width = (box.width - 0.4) / 2.0
                q_height = 0.55
                for i, step in enumerate(steps[:4]):
                    col = i % 2
                    row = i // 2
                    c_left = box.left + col * (q_width + 0.4)
                    c_top = q_top + row * (q_height + 0.1)
                    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(c_top), Inches(q_width), Inches(q_height))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = palette["accent"]
                    shape.line.color.rgb = palette["text"]
                    shape.line.width = Pt(1)
                    tf = shape.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = f"Q{i + 1}: {step}"
                    set_run_style(run, font_size=10, bold=True, color=palette["background"])

            elif diag_type == "cycle":
                import math
                center_x = box.left + (box.width / 2.0)
                center_y = box.top + 0.95
                radius_x = min(2.8, (box.width - 2.5) / 2.0)
                radius_y = 0.50
                num_s = len(steps)

                center_badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center_x - 0.45), Inches(center_y - 0.35), Inches(0.9), Inches(0.7))
                center_badge.fill.solid()
                center_badge.fill.fore_color.rgb = palette["accent"]
                tf_c = center_badge.text_frame
                tf_c.word_wrap = True
                tf_c.clear()
                p_c = tf_c.paragraphs[0]
                p_c.alignment = PP_ALIGN.CENTER
                run_c = p_c.add_run()
                run_c.text = "🔁 LOOP"
                set_run_style(run_c, font_size=10, bold=True, color=palette["background"])

                for i, step in enumerate(steps):
                    angle = (2.0 * math.pi * i / num_s) - (math.pi / 2.0)
                    c_x = center_x + radius_x * math.cos(angle)
                    c_y = center_y + radius_y * math.sin(angle)
                    card_w = 1.6
                    card_h = 0.45
                    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_x - (card_w / 2.0)), Inches(c_y - (card_h / 2.0)), Inches(card_w), Inches(card_h))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = palette["background"]
                    shape.line.color.rgb = palette["accent"]
                    shape.line.width = Pt(2)
                    tf = shape.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = f"{i + 1}. {step}"
                    set_run_style(run, font_size=10, bold=True, color=palette["text"])

            elif diag_type == "timeline":
                t_line_y = box.top + 0.9
                axis = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.left + 0.5), Inches(t_line_y), Inches(box.width - 1.0), Inches(0.06))
                axis.fill.solid()
                axis.fill.fore_color.rgb = palette["accent"]
                axis.line.fill.background()

                card_width = (box.width - 0.5) / len(steps)
                for i, step in enumerate(steps):
                    c_x = box.left + 0.25 + i * card_width
                    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_x + (card_width / 2.0) - 0.15), Inches(t_line_y - 0.12), Inches(0.3), Inches(0.3))
                    dot.fill.solid()
                    dot.fill.fore_color.rgb = palette["accent"]

                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_x + 0.05), Inches(t_line_y + 0.25), Inches(card_width - 0.1), Inches(0.55))
                    card.fill.solid()
                    card.fill.fore_color.rgb = palette["background"]
                    card.line.color.rgb = palette["accent"]
                    card.line.width = Pt(1)
                    tf = card.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = f"M{i + 1}: {step}"
                    set_run_style(run, font_size=10, bold=True, color=palette["text"])

            elif diag_type == "io_cards":
                c_top = box.top + 0.4
                c_width = (box.width - 0.6) / max(3, len(steps))
                for i, step in enumerate(steps):
                    c_left = box.left + i * (c_width + 0.2)
                    labels = ["INPUT DATA", "PROCESSING", "OUTPUT RESULT"]
                    label_title = labels[i] if i < len(labels) else f"STEP {i + 1}"
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(c_top), Inches(c_width), Inches(0.75))
                    card.fill.solid()
                    card.fill.fore_color.rgb = palette["accent"]
                    card.line.color.rgb = palette["text"]
                    tf = card.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p0 = tf.paragraphs[0]
                    p0.alignment = PP_ALIGN.CENTER
                    r0 = p0.add_run()
                    r0.text = f"[{label_title}]\n"
                    set_run_style(r0, font_size=9, bold=True, color=palette["background"])
                    r1 = p0.add_run()
                    r1.text = step
                    set_run_style(r1, font_size=11, bold=True, color=palette["background"])

            elif diag_type == "mindmap":
                c_top = box.top + 0.4
                center_card = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(box.left + (box.width / 2.0) - 1.25), Inches(c_top), Inches(2.5), Inches(0.45))
                center_card.fill.solid()
                center_card.fill.fore_color.rgb = palette["accent"]
                tf_c = center_card.text_frame
                tf_c.word_wrap = True
                tf_c.clear()
                p_c = tf_c.paragraphs[0]
                p_c.alignment = PP_ALIGN.CENTER
                run_c = p_c.add_run()
                run_c.text = f"🧠 {steps[0]}"
                set_run_style(run_c, font_size=11, bold=True, color=palette["background"])

                sub_steps = steps[1:]
                if sub_steps:
                    sub_top = c_top + 0.6
                    sub_w = (box.width - 0.4) / max(1, len(sub_steps))
                    for i, s_step in enumerate(sub_steps):
                        s_left = box.left + i * sub_w
                        scard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(s_left + 0.05), Inches(sub_top), Inches(sub_w - 0.1), Inches(0.45))
                        scard.fill.solid()
                        scard.fill.fore_color.rgb = palette["background"]
                        scard.line.color.rgb = palette["accent"]
                        tf_s = scard.text_frame
                        tf_s.word_wrap = True
                        tf_s.clear()
                        p_s = tf_s.paragraphs[0]
                        p_s.alignment = PP_ALIGN.CENTER
                        run_s = p_s.add_run()
                        run_s.text = f"🔹 {s_step}"
                        set_run_style(run_s, font_size=10, bold=True, color=palette["text"])

            elif diag_type == "tree":
                c_top = box.top + 0.35
                root_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(box.left + (box.width / 2.0) - 1.5), Inches(c_top), Inches(3.0), Inches(0.45))
                root_card.fill.solid()
                root_card.fill.fore_color.rgb = palette["accent"]
                root_card.line.color.rgb = palette["text"]
                tf_r = root_card.text_frame
                tf_r.word_wrap = True
                tf_r.clear()
                p_r = tf_r.paragraphs[0]
                p_r.alignment = PP_ALIGN.CENTER
                run_r = p_r.add_run()
                run_r.text = f"🌳 {steps[0]}"
                set_run_style(run_r, font_size=11, bold=True, color=palette["background"])

                branches = steps[1:]
                if branches:
                    branch_top = c_top + 0.65
                    b_w = (box.width - 0.4) / max(1, len(branches))
                    for i, b_step in enumerate(branches):
                        b_left = box.left + i * b_w
                        bcard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(b_left + 0.05), Inches(branch_top), Inches(b_w - 0.1), Inches(0.45))
                        bcard.fill.solid()
                        bcard.fill.fore_color.rgb = palette["background"]
                        bcard.line.color.rgb = palette["accent"]
                        bcard.line.width = Pt(1.5)
                        tf_b = bcard.text_frame
                        tf_b.word_wrap = True
                        tf_b.clear()
                        p_b = tf_b.paragraphs[0]
                        p_b.alignment = PP_ALIGN.CENTER
                        run_b = p_b.add_run()
                        run_b.text = f"🌿 {b_step}"
                        set_run_style(run_b, font_size=10, bold=True, color=palette["text"])

            else:
                card_top = box.top + 0.4
                card_height = 0.65
                total_width = box.width
                num_steps = len(steps)
                gap = 0.20
                card_width = max(1.4, (total_width - (gap * (num_steps - 1))) / num_steps)

                for i, step in enumerate(steps):
                    c_left = box.left + i * (card_width + gap)
                    is_terminal = (i == 0 or i == num_steps - 1) and diag_type == "flowchart"
                    shape_type = MSO_SHAPE.OVAL if is_terminal else MSO_SHAPE.ROUNDED_RECTANGLE
                    shape = slide.shapes.add_shape(shape_type, Inches(c_left), Inches(card_top), Inches(card_width), Inches(card_height))
                    try:
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = palette["accent"]
                        shape.line.color.rgb = palette["text"]
                        shape.line.width = Pt(1)
                    except Exception:
                        pass

                    tf = shape.text_frame
                    tf.word_wrap = True
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = step
                    set_run_style(run, font_size=11, bold=True, color=palette["background"])

                    if i < num_steps - 1:
                        arrow_box = slide.shapes.add_textbox(Inches(c_left + card_width), Inches(card_top + 0.15), Inches(gap), Inches(0.35))
                        ap = arrow_box.text_frame.paragraphs[0]
                        ap.alignment = PP_ALIGN.CENTER
                        ap.text = "⚔️" if diag_type == "comparison" else "➔"
                        ap.font.size = Pt(13)
                        ap.font.color.rgb = palette["accent"]
        else:
            diag_box = slide.shapes.add_textbox(Inches(box.left), Inches(box.top + 0.35), Inches(box.width), Inches(box.height - 0.35))
            tf = diag_box.text_frame
            tf.clear()
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = diagram_text
            set_run_style(run, font_size=14, bold=True, color=palette["text"])

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        theme = plan.get("theme_name")
        steps = safe_list(plan.get("steps")) or safe_list(plan.get("points"))
        diag_type = str(plan.get("diag_type", plan.get("diagram_type", "stack"))).lower()

        if diag_type in {"stack", "vertical_stack", "architecture_stack", "funnel", "pyramid"}:
            needed_height = max(1.8, 0.45 + len(steps) * 0.52)
        elif diag_type in {"cycle", "mindmap", "timeline"}:
            needed_height = 2.4
        else:
            needed_height = 1.8

        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": needed_height}}, theme_name=theme)
        return current_y + needed_height + 0.35


class ImagePlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        url = normalize_whitespace(plan.get("url", ""))
        path = normalize_whitespace(plan.get("path", ""))
        caption = normalize_whitespace(plan.get("caption", ""))
        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.8, top_pos, 6.6, 3.5))

        safe_top = min(raw_box.top, 4.2)
        caption_space = 0.35

        custom_h = plan.get("img_height") or plan.get("height") or plan.get("size")
        if custom_h and str(custom_h).replace(".", "", 1).isdigit() and float(custom_h) > 0:
            custom_height_in_inches = (float(custom_h) / 180.0) * 3.2
            safe_height = max(1.0, min(4.2, round(custom_height_in_inches, 2)))
        else:
            safe_height = min(raw_box.height, max(1.5, round(5.8 - safe_top - caption_space, 2)))

        box = Box(raw_box.left, safe_top, raw_box.width, safe_height)

        target_source = url or path
        safe_path = sanitize_image_path(target_source)
        if not safe_path:
            raw_q = caption or plan.get("title") or "presentation visual"
            pres_title = plan.get("presentation_title") or plan.get("slide_title") or ""
            query = f"{pres_title} {raw_q}".strip()
            fetched = fetch_unsplash_url(query) or fetch_unsplash_image_for_topic(query) or fetch_unsplash_url(raw_q)
            if fetched:
                safe_path = sanitize_image_path(fetched)

        if safe_path:
            try:
                from PIL import Image as PILImage
                with PILImage.open(safe_path) as img:
                    img_w, img_h = img.size

                aspect = img_w / img_h if img_h > 0 else 1.0
                target_aspect = box.width / box.height if box.height > 0 else 1.0

                if aspect > target_aspect:
                    render_w = box.width
                    render_h = box.width / aspect
                else:
                    render_h = box.height
                    render_w = box.height * aspect

                render_h = min(render_h, max(1.0, 5.8 - box.top))
                pos_left = box.left + (box.width - render_w) / 2
                pos_top = min(5.6, box.top + (box.height - render_h) / 2)
                pos_top = max(1.5, pos_top)
                if pos_top + render_h > 5.9:
                    render_h = max(1.0, 5.9 - pos_top)

                slide.shapes.add_picture(
                    safe_path,
                    Inches(pos_left),
                    Inches(pos_top),
                    width=Inches(render_w),
                    height=Inches(render_h),
                )

                display_label = caption or plan.get("title") or "Visual"
                cap_top = min(6.0, pos_top + render_h + 0.05)
                cap_box = slide.shapes.add_textbox(Inches(pos_left), Inches(cap_top), Inches(render_w), Inches(0.32))
                cap_tf = cap_box.text_frame
                cap_tf.word_wrap = True
                p = cap_tf.paragraphs[0]
                p.text = f"fig:- {display_label}"
                p.alignment = PP_ALIGN.CENTER
                set_run_style(p.runs[0] if p.runs else p.add_run(), font_size=10, bold=True, color=palette["accent"])
                return
            except Exception as exc:
                logger.warning("Failed to insert picture %s: %s", safe_path, exc)
                safe_path = None

        if not safe_path:
            box_shape = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
            tf = box_shape.text_frame
            tf.text = f"Visual: {caption or path or 'Topic'}"
            tf.paragraphs[0].font.size = Pt(18)
            try:
                tf.paragraphs[0].runs[0].font.color.rgb = palette["text"]
            except Exception:
                pass

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        avail_h = max(1.5, round(6.5 - current_y - 0.35, 2))
        img_h = min(3.4, avail_h)
        img_w = min(7.5, content_width)
        img_left = (13.333 - img_w) / 2.0
        self.apply(slide, {**plan, "top": current_y, "box": {"left": img_left, "top": current_y, "width": img_w, "height": img_h}}, theme_name=None)
        return current_y + img_h + 0.45


class TablePlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        effective_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(effective_theme)
        headers = safe_list(plan.get("headers"))
        rows = safe_list(plan.get("rows"))
        top_pos = float(plan.get("top", 1.6))
        raw_box = as_box(plan, Box(0.8, top_pos, 11.7, 3.5))

        safe_top = min(raw_box.top, 4.2)
        safe_height = min(raw_box.height, max(1.5, round(6.6 - safe_top, 2)))
        box = Box(raw_box.left, safe_top, raw_box.width, safe_height)

        if not headers or not rows:
            box_shape = slide.shapes.add_textbox(Inches(1.0), Inches(safe_top), Inches(8.0), Inches(1.0))
            tf = box_shape.text_frame
            tf.text = "Table data not found or incomplete."
            tf.paragraphs[0].font.size = Pt(18)
            return

        slide_title = str(plan.get("slide_title") or plan.get("title") or "").lower()
        if re.search(r"\b(?:vs\.?|versus|compared\s+to)\b", slide_title):
            if len(headers) > 3:
                headers = headers[:3]
                rows = [r[:3] if isinstance(r, (list, tuple)) else r for r in rows]

        custom_header_bg = plan.get("header_bg")
        custom_header_color = plan.get("header_color")
        custom_cell_bg = plan.get("cell_bg")
        custom_cell_color = plan.get("cell_color")

        raw_h_fs = plan.get("header_font_size")
        font_size_header = int(raw_h_fs) if raw_h_fs and str(raw_h_fs).isdigit() else 13

        raw_c_fs = plan.get("cell_font_size", plan.get("font_size"))
        font_size_cell = int(raw_c_fs) if raw_c_fs and str(raw_c_fs).isdigit() else 11

        align_opt = str(plan.get("align", plan.get("alignment", "left"))).lower()
        align_map = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY, "left": PP_ALIGN.LEFT}
        cell_align = align_map.get(align_opt, PP_ALIGN.LEFT)

        bg_is_light = is_light_color(palette["background"])

        if custom_header_bg:
            hdr_bg_rgb = hex_to_rgb(str(custom_header_bg))
        else:
            hdr_bg_rgb = palette.get("table_header_bg") or palette["accent"]

        if custom_header_color:
            hdr_txt_rgb = hex_to_rgb(str(custom_header_color))
        else:
            hdr_txt_rgb = palette.get("table_header_text") or (RGBColor(15, 23, 42) if is_light_color(hdr_bg_rgb) else RGBColor(255, 255, 255))
        hdr_txt_rgb = ensure_readable_text_color(hdr_bg_rgb, hdr_txt_rgb)

        if custom_cell_bg:
            base_row_rgb1 = hex_to_rgb(str(custom_cell_bg))
            base_row_rgb2 = base_row_rgb1
        else:
            base_row_rgb1 = palette.get("table_row_bg1") or (RGBColor(241, 245, 249) if bg_is_light else RGBColor(30, 41, 59))
            base_row_rgb2 = palette.get("table_row_bg2") or (RGBColor(255, 255, 255) if bg_is_light else palette["background"])

        default_row_txt = palette.get("table_row_text") or palette["text"]

        cols = len(headers)
        row_count = len(rows) + 1
        table_shape = slide.shapes.add_table(row_count, cols, Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
        table = table_shape.table

        for c, header in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = str(header)
            try:
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            except Exception:
                pass
            try:
                cell.fill.solid()
                cell.fill.fore_color.rgb = hdr_bg_rgb
            except Exception:
                pass
            tf = cell.text_frame
            tf.word_wrap = True
            try:
                tf.margin_left = Inches(0.08)
                tf.margin_right = Inches(0.08)
                tf.margin_top = Inches(0.06)
                tf.margin_bottom = Inches(0.06)
            except Exception:
                pass
            for p in tf.paragraphs:
                p.alignment = cell_align
                for run in p.runs:
                    set_run_style(run, font_size=font_size_header, bold=True, color=hdr_txt_rgb)

        for r, row in enumerate(rows, start=1):
            row_bg_rgb = base_row_rgb1 if r % 2 == 1 else base_row_rgb2
            if custom_cell_color:
                row_txt_rgb = hex_to_rgb(str(custom_cell_color))
            else:
                row_txt_rgb = ensure_readable_text_color(row_bg_rgb, default_row_txt)

            if isinstance(row, (list, tuple)):
                row_list = list(row)
            elif hasattr(row, "__iter__") and not isinstance(row, (str, bytes, dict)):
                row_list = list(row)
            else:
                row_list = [str(row)]

            values = row_list + [""] * max(0, cols - len(row_list))
            values = values[:cols]
            for c, value in enumerate(values):
                cell = table.cell(r, c)
                cell.text = str(value)
                try:
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                except Exception:
                    pass
                try:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = row_bg_rgb
                except Exception:
                    pass
                tf = cell.text_frame
                tf.word_wrap = True
                try:
                    tf.margin_left = Inches(0.08)
                    tf.margin_right = Inches(0.08)
                    tf.margin_top = Inches(0.06)
                    tf.margin_bottom = Inches(0.06)
                except Exception:
                    pass
                for p in tf.paragraphs:
                    p.alignment = cell_align
                    for run in p.runs:
                        set_run_style(run, font_size=font_size_cell, color=row_txt_rgb)

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        avail_h = max(1.5, round(6.5 - current_y, 2))
        tbl_h = min(3.4, avail_h)
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": tbl_h}}, theme_name=eff_theme)
        return current_y + tbl_h + 0.35


class NotesPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        set_slide_notes(slide, plan.get("notes", ""))

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        set_slide_notes(slide, plan.get("notes", ""))
        return current_y


class StatPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        number = str(plan.get("number", "100%")).strip()
        label = str(plan.get("label", "Metric")).strip()
        top_pos = float(plan.get("top", 1.5))
        box = as_box(plan, Box(0.8, top_pos, 11.7, 1.2))

        s_box = slide.shapes.add_textbox(Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
        tf = s_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r_num = p.add_run()
        r_num.text = f"{number} "
        set_run_style(r_num, font_size=40, bold=True, color=palette["accent"])
        r_lbl = p.add_run()
        r_lbl.text = label
        set_run_style(r_lbl, font_size=18, bold=False, color=palette["text"])

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        *,
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        number = str(plan.get("number", "100%")).strip()
        label = str(plan.get("label", "Metric")).strip()
        box_height = 0.8
        s_box = slide.shapes.add_textbox(Inches(left_margin), Inches(current_y), Inches(content_width), Inches(box_height))
        tf = s_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r_num = p.add_run()
        r_num.text = f"{number} "
        set_run_style(r_num, font_size=36, bold=True, color=palette["accent"])
        r_lbl = p.add_run()
        r_lbl.text = label
        set_run_style(r_lbl, font_size=16, bold=False, color=palette["text"])
        return current_y + box_height + 0.15


class Paragraph2ColPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        palette = get_theme_palette(theme_name)
        raw_items = plan.get("items") or plan.get("paragraphs") or plan.get("columns")
        items: List[Dict[str, str]] = []

        if isinstance(raw_items, list) and len(raw_items) > 0:
            for it in raw_items:
                if isinstance(it, dict):
                    t = str(it.get("title") or "").strip()
                    txt = str(it.get("text") or it.get("paragraph") or "").strip()
                    if t or txt:
                        items.append({"title": t, "text": txt})
                elif isinstance(it, str) and it.strip():
                    items.append({"title": "", "text": it.strip()})

        if not items:
            left_text = str(plan.get("left_text") or plan.get("text") or "").strip()
            right_text = str(plan.get("right_text") or "").strip()
            left_title = str(plan.get("left_title") or "").strip()
            right_title = str(plan.get("right_title") or "").strip()

            if left_text or left_title:
                items.append({"title": left_title, "text": left_text})
            if right_text or right_title:
                items.append({"title": right_title, "text": right_text})

        if not items:
            return

        top_pos = float(plan.get("top", 2.0))
        raw_box = as_box(plan, Box(0.9, top_pos, 11.5, 3.8))
        box = Box(raw_box.left, raw_box.top, raw_box.width, max(1.5, raw_box.height))

        num_items = len(items)
        cols_per_row = num_items if num_items <= 4 else 3
        num_rows = (num_items + cols_per_row - 1) // cols_per_row
        row_height = 2.4 if num_rows == 1 else 1.8
        gap = 0.35

        for i, item in enumerate(items):
            row_idx = i // cols_per_row
            col_idx = i % cols_per_row
            items_in_this_row = min(cols_per_row, num_items - row_idx * cols_per_row)

            w = (box.width - (gap * (items_in_this_row - 1))) / items_in_this_row
            left = box.left + col_idx * (w + gap)
            top = box.top + row_idx * (row_height + 0.2)

            item_box = Box(left, top, w, row_height)
            add_card_container(slide, item_box, palette)

            tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(row_height))
            tf = tb.text_frame
            tf.clear()
            tf.word_wrap = True

            if item.get("title"):
                p_t = tf.paragraphs[0]
                run_t = p_t.add_run()
                run_t.text = f"{item['title']}\n"
                set_run_style(run_t, font_size=12 if num_items > 3 else 13, bold=True, color=palette["accent"])

            pts = item.get("points") if isinstance(item, dict) else None
            raw_text = str(item.get("text") if isinstance(item, dict) else item or "").strip()
            if not pts and raw_text and ("\n" in raw_text or any(raw_text.startswith(p) for p in ("•", "-", "*", "1.", "2.", "✓", "➔"))):
                lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
                if len(lines) > 1 or (lines and any(lines[0].startswith(p) for p in ("•", "-", "*", "1.", "2.", "✓", "➔"))):
                    pts = [re.sub(r"^[•\-*✓➔\d+\.\s]+", "", l).strip() for l in lines]

            if pts and isinstance(pts, list) and len(pts) > 0:
                b_style = (item.get("bullet_style") if isinstance(item, dict) else None) or plan.get("bullet_style") or "disc"
                for idx_pt, pt in enumerate(pts):
                    p_pt = tf.add_paragraph()
                    prefix = format_bullet_prefix(b_style, idx_pt, pts)
                    p_pt.text = f"{prefix} {pt}"
                    p_pt.space_after = Pt(2)
                    run_pt = p_pt.runs[0] if p_pt.runs else p_pt.add_run()
                    set_run_style(run_pt, font_size=10 if num_items > 3 else 11, color=palette["text"])
            elif item.get("title"):
                p_txt = tf.add_paragraph()
                run_txt = p_txt.add_run()
                run_txt.text = raw_text
                set_run_style(run_txt, font_size=10 if num_items > 3 else 11, color=palette["text"])
            else:
                tf.text = raw_text
                configure_text_frame(tf, font_size=11, color=palette["text"])

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
    ) -> float:
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 3.0}}, theme_name=None)
        return current_y + 3.2
class CalloutPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        eff_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(eff_theme)
        text = str(plan.get("text") or plan.get("takeaway") or plan.get("quote") or "").strip()
        title = str(plan.get("title") or plan.get("header") or "KEY TAKEAWAY").strip()
        icon = str(plan.get("icon") or "💡").strip()
        if not text:
            return

        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.9, top_pos, 11.5, 1.4))
        box = Box(raw_box.left, min(raw_box.top, 4.5), raw_box.width, max(1.2, raw_box.height))

        card = add_card_container(slide, box, palette)
        try:
            accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.left), Inches(box.top), Inches(0.12), Inches(box.height))
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = palette["accent"]
            accent_bar.line.fill.background()
        except Exception:
            pass

        tbox = slide.shapes.add_textbox(Inches(box.left + 0.25), Inches(box.top + 0.1), Inches(box.width - 0.4), Inches(box.height - 0.2))
        tf = tbox.text_frame
        tf.clear()
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = f"{icon} {title.upper()}\n"
        set_run_style(r0, font_size=12, bold=True, color=palette["accent"])

        p1 = tf.add_paragraph()
        p1.space_before = Pt(4)
        r1 = p1.add_run()
        r1.text = f'"{text}"' if "quote" in str(plan.get("type", "")).lower() else text
        set_run_style(r1, font_size=13, color=palette["text"])

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 1.4}}, theme_name=eff_theme)
        return current_y + 1.6


class KPIGridPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        eff_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(eff_theme)
        kpis = safe_list(plan.get("kpis") or plan.get("items") or plan.get("stats"))
        if not kpis:
            return

        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.8, top_pos, 11.7, 1.8))
        box = Box(raw_box.left, min(raw_box.top, 4.2), raw_box.width, max(1.5, raw_box.height))

        num_cards = min(4, len(kpis))
        gap = 0.25
        card_w = (box.width - (gap * (num_cards - 1))) / num_cards

        for i, kpi in enumerate(kpis[:num_cards]):
            if isinstance(kpi, dict):
                number = str(kpi.get("number") or kpi.get("value") or "100%").strip()
                label = str(kpi.get("label") or kpi.get("title") or "Metric").strip()
                trend = str(kpi.get("trend") or kpi.get("change") or "").strip()
            else:
                number = str(kpi).strip()
                label = f"Metric {i + 1}"
                trend = ""

            c_left = box.left + i * (card_w + gap)
            card_box = Box(c_left, box.top, card_w, box.height)
            add_card_container(slide, card_box, palette)

            tbox = slide.shapes.add_textbox(Inches(c_left + 0.1), Inches(box.top + 0.15), Inches(card_w - 0.2), Inches(box.height - 0.3))
            tf = tbox.text_frame
            tf.clear()
            tf.word_wrap = True

            p0 = tf.paragraphs[0]
            p0.alignment = PP_ALIGN.CENTER
            r0 = p0.add_run()
            r0.text = f"{number}\n"
            set_run_style(r0, font_size=30, bold=True, color=palette["accent"])

            p1 = tf.add_paragraph()
            p1.alignment = PP_ALIGN.CENTER
            p1.space_before = Pt(2)
            r1 = p1.add_run()
            r1.text = label
            set_run_style(r1, font_size=11, bold=True, color=palette["text"])

            if trend:
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                p2.space_before = Pt(4)
                r2 = p2.add_run()
                r2.text = f" {trend} "
                trend_color = RGBColor(16, 185, 129) if "+" in trend or "↗" in trend else palette["accent"]
                set_run_style(r2, font_size=10, bold=True, color=trend_color)

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 1.8}}, theme_name=eff_theme)
        return current_y + 2.05


class ProsConsPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        eff_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(eff_theme)
        pros = safe_list(plan.get("pros") or plan.get("strengths") or plan.get("advantages"))
        cons = safe_list(plan.get("cons") or plan.get("weaknesses") or plan.get("challenges"))
        pros_title = str(plan.get("pros_title") or "✅ STRENGTHS & ADVANTAGES").strip()
        cons_title = str(plan.get("cons_title") or "❌ CHALLENGES & CONSIDERATIONS").strip()

        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.8, top_pos, 11.7, 3.2))
        box = Box(raw_box.left, min(raw_box.top, 4.2), raw_box.width, max(2.2, raw_box.height))

        card_w = (box.width - 0.4) / 2.0

        left_box = Box(box.left, box.top, card_w, box.height)
        add_card_container(slide, left_box, palette, border_color=RGBColor(16, 185, 129))

        tb_pros = slide.shapes.add_textbox(Inches(left_box.left + 0.15), Inches(left_box.top + 0.15), Inches(card_w - 0.3), Inches(box.height - 0.3))
        tf_p = tb_pros.text_frame
        tf_p.clear()
        tf_p.word_wrap = True
        p_hdr = tf_p.paragraphs[0]
        r_hdr = p_hdr.add_run()
        r_hdr.text = f"{pros_title}\n"
        set_run_style(r_hdr, font_size=13, bold=True, color=RGBColor(16, 185, 129))

        for item in pros:
            p = tf_p.add_paragraph()
            p.space_before = Pt(4)
            r = p.add_run()
            r.text = f"✔ {item}"
            set_run_style(r, font_size=11, color=palette["text"])

        right_box = Box(box.left + card_w + 0.4, box.top, card_w, box.height)
        add_card_container(slide, right_box, palette, border_color=RGBColor(239, 68, 68))

        tb_cons = slide.shapes.add_textbox(Inches(right_box.left + 0.15), Inches(right_box.top + 0.15), Inches(card_w - 0.3), Inches(box.height - 0.3))
        tf_c = tb_cons.text_frame
        tf_c.clear()
        tf_c.word_wrap = True
        p_chdr = tf_c.paragraphs[0]
        r_chdr = p_chdr.add_run()
        r_chdr.text = f"{cons_title}\n"
        set_run_style(r_chdr, font_size=13, bold=True, color=RGBColor(239, 68, 68))

        for item in cons:
            p = tf_c.add_paragraph()
            p.space_before = Pt(4)
            r = p.add_run()
            r.text = f"✖ {item}"
            set_run_style(r, font_size=11, color=palette["text"])

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 3.2}}, theme_name=eff_theme)
        return current_y + 3.45


class RoadmapPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        eff_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(eff_theme)
        phases = safe_list(plan.get("phases") or plan.get("milestones") or plan.get("steps"))
        if not phases:
            phases = [
                {"phase": "Q1 2026", "title": "Architecture & Specs", "status": "COMPLETED"},
                {"phase": "Q2 2026", "title": "Core Platform Build", "status": "IN PROGRESS"},
                {"phase": "Q3 2026", "title": "Market Integration", "status": "PLANNED"},
                {"phase": "Q4 2026", "title": "Global Scaling", "status": "PLANNED"},
            ]

        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.8, top_pos, 11.7, 2.2))
        box = Box(raw_box.left, min(raw_box.top, 4.2), raw_box.width, max(1.8, raw_box.height))

        num_p = min(4, len(phases))
        gap = 0.25
        card_w = (box.width - (gap * (num_p - 1))) / num_p

        axis_y = box.top + 0.3
        try:
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.left + 0.2), Inches(axis_y), Inches(box.width - 0.4), Inches(0.05))
            line.fill.solid()
            line.fill.fore_color.rgb = palette["accent"]
            line.line.fill.background()
        except Exception:
            pass

        for i, item in enumerate(phases[:num_p]):
            if isinstance(item, dict):
                phase_hdr = str(item.get("phase") or item.get("quarter") or f"PHASE {i + 1}").strip()
                title = str(item.get("title") or item.get("label") or f"Milestone {i + 1}").strip()
                status = str(item.get("status") or "PLANNED").upper().strip()
            else:
                phase_hdr = f"PHASE {i + 1}"
                title = str(item).strip()
                status = "PLANNED"

            c_left = box.left + i * (card_w + gap)
            try:
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c_left + (card_w / 2.0) - 0.15), Inches(axis_y - 0.12), Inches(0.3), Inches(0.3))
                dot.fill.solid()
                dot.fill.fore_color.rgb = palette["accent"]
                dot.line.color.rgb = palette["text"]
            except Exception:
                pass

            card_box = Box(c_left, axis_y + 0.3, card_w, box.height - 0.5)
            add_card_container(slide, card_box, palette)

            tbox = slide.shapes.add_textbox(Inches(c_left + 0.1), Inches(card_box.top + 0.1), Inches(card_w - 0.2), Inches(card_box.height - 0.2))
            tf = tbox.text_frame
            tf.clear()
            tf.word_wrap = True

            p0 = tf.paragraphs[0]
            p0.alignment = PP_ALIGN.CENTER
            r0 = p0.add_run()
            r0.text = f"{phase_hdr}\n"
            set_run_style(r0, font_size=11, bold=True, color=palette["accent"])

            p1 = tf.add_paragraph()
            p1.alignment = PP_ALIGN.CENTER
            p1.space_before = Pt(2)
            r1 = p1.add_run()
            r1.text = f"{title}\n"
            set_run_style(r1, font_size=11, bold=False, color=palette["text"])

            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.space_before = Pt(4)
            r2 = p2.add_run()
            r2.text = f" [{status}] "
            st_color = RGBColor(16, 185, 129) if "COMPLET" in status else (RGBColor(59, 130, 246) if "PROGRESS" in status else RGBColor(148, 163, 184))
            set_run_style(r2, font_size=9, bold=True, color=st_color)

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 2.2}}, theme_name=eff_theme)
        return current_y + 2.45


class CodeBlockPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        eff_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(eff_theme)
        code = str(plan.get("code") or plan.get("snippet") or 'def main():\n    print("Hello Antigravity Engine")').strip()
        title = str(plan.get("title") or plan.get("filename") or "code_snippet.py").strip()
        lang = str(plan.get("language") or plan.get("lang") or "python").upper().strip()

        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.9, top_pos, 11.5, 3.2))
        box = Box(raw_box.left, min(raw_box.top, 4.0), raw_box.width, max(2.0, raw_box.height))

        try:
            terminal_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(box.left), Inches(box.top), Inches(box.width), Inches(box.height))
            terminal_bg.fill.solid()
            terminal_bg.fill.fore_color.rgb = RGBColor(15, 23, 42)
            terminal_bg.line.color.rgb = RGBColor(51, 65, 85)
            terminal_bg.line.width = Pt(1.5)
        except Exception:
            pass

        try:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.left), Inches(box.top), Inches(box.width), Inches(0.35))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(30, 41, 59)
            bar.line.fill.background()

            bar_tbox = slide.shapes.add_textbox(Inches(box.left + 0.15), Inches(box.top + 0.02), Inches(box.width - 0.3), Inches(0.3))
            tf_b = bar_tbox.text_frame
            tf_b.clear()
            p_b = tf_b.paragraphs[0]
            r_dots = p_b.add_run()
            r_dots.text = "🔴 🟡 🟢   "
            set_run_style(r_dots, font_size=9)
            r_title = p_b.add_run()
            r_title.text = f"{title} ({lang})"
            set_run_style(r_title, font_size=10, bold=True, color=RGBColor(148, 163, 184))
        except Exception:
            pass

        code_tbox = slide.shapes.add_textbox(Inches(box.left + 0.2), Inches(box.top + 0.4), Inches(box.width - 0.4), Inches(box.height - 0.45))
        tf_c = code_tbox.text_frame
        tf_c.clear()
        tf_c.word_wrap = True
        
        for line in code.splitlines():
            p = tf_c.add_paragraph() if tf_c.paragraphs[0].text else tf_c.paragraphs[0]
            r = p.add_run()
            r.text = line
            r.font.name = "Consolas"
            set_run_style(r, font_size=11, color=RGBColor(248, 250, 252))

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 3.2}}, theme_name=eff_theme)
        return current_y + 3.45


class SpeakerCardPlugin(BasePlugin):
    def apply(self, slide, plan: Dict[str, Any], theme_name: Optional[str] = None) -> None:
        eff_theme = theme_name or plan.get("theme_name")
        palette = get_theme_palette(eff_theme)
        name = str(plan.get("name") or plan.get("speaker") or plan.get("presenter") or "Speaker Name").strip()
        role = str(plan.get("role") or plan.get("title") or "Keynote Presenter").strip()
        bio_points = safe_list(plan.get("bio") or plan.get("highlights") or [f"Executive Lead at {name}", "Domain Expert & Keynote Presenter"])

        top_pos = float(plan.get("top", 1.8))
        raw_box = as_box(plan, Box(0.9, top_pos, 11.5, 2.2))
        box = Box(raw_box.left, min(raw_box.top, 4.2), raw_box.width, max(1.8, raw_box.height))

        add_card_container(slide, box, palette)

        try:
            avatar = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(box.left + 0.3), Inches(box.top + 0.3), Inches(1.2), Inches(1.2))
            avatar.fill.solid()
            avatar.fill.fore_color.rgb = palette["accent"]
            avatar.line.color.rgb = palette["text"]
            tf_a = avatar.text_frame
            tf_a.clear()
            p_a = tf_a.paragraphs[0]
            p_a.alignment = PP_ALIGN.CENTER
            r_a = p_a.add_run()
            r_a.text = "👤"
            set_run_style(r_a, font_size=28)
        except Exception:
            pass

        tbox = slide.shapes.add_textbox(Inches(box.left + 1.7), Inches(box.top + 0.2), Inches(box.width - 1.9), Inches(box.height - 0.4))
        tf = tbox.text_frame
        tf.clear()
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = f"{name}\n"
        set_run_style(r0, font_size=18, bold=True, color=palette["accent"])

        p1 = tf.add_paragraph()
        r1 = p1.add_run()
        r1.text = f"{role}\n"
        set_run_style(r1, font_size=12, bold=True, color=palette["text"])

        for pt in bio_points:
            p = tf.add_paragraph()
            p.space_before = Pt(3)
            r = p.add_run()
            r.text = f"• {pt}"
            set_run_style(r, font_size=11, color=palette["text"])

    def apply_with_y(
        self,
        slide,
        plan: Dict[str, Any],
        current_y: float,
        left_margin: float,
        content_width: float,
        palette: Dict[str, RGBColor],
        theme_name: Optional[str] = None,
    ) -> float:
        eff_theme = theme_name or plan.get("theme_name")
        self.apply(slide, {**plan, "top": current_y, "box": {"left": left_margin, "top": current_y, "width": content_width, "height": 2.2}}, theme_name=eff_theme)
        return current_y + 2.45


PLUGIN_REGISTRY: Dict[str, BasePlugin] = {
    "text": TextPlugin(),
    "paragraph": ParagraphPlugin(),
    "paragraph_2col": Paragraph2ColPlugin(),
    "bullets": BulletsPlugin(),
    "chart": ChartPlugin(),
    "image": ImagePlugin(),
    "table": TablePlugin(),
    "notes": NotesPlugin(),
    "diagram": DiagramPlugin(),
    "stat": StatPlugin(),
    "callout": CalloutPlugin(),
    "kpi_grid": KPIGridPlugin(),
    "pros_cons": ProsConsPlugin(),
    "roadmap": RoadmapPlugin(),
    "code_block": CodeBlockPlugin(),
    "speaker_card": SpeakerCardPlugin(),
}


def add_brand_elements_to_slide(slide, plan: PresentationPlan, slide_width_in: float, is_cover: bool):
    brand_footer = getattr(plan, "brand_footer", None)
    if brand_footer and str(brand_footer).strip():
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(5.5), Inches(0.3))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = str(brand_footer).strip()
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = RGBColor(148, 163, 184)
        brand_font = getattr(plan, "brand_font", None)
        if brand_font:
            p.font.name = str(brand_font).strip()

    brand_logo = getattr(plan, "brand_logo", None)
    if brand_logo and isinstance(brand_logo, str) and len(brand_logo.strip()) > 5:
        try:
            import base64, io, urllib.request
            image_stream = None
            logo_str = brand_logo.strip()
            if logo_str.startswith("data:image/"):
                header, encoded = logo_str.split(",", 1)
                data = base64.b64decode(encoded)
                image_stream = io.BytesIO(data)
            elif logo_str.startswith("http://") or logo_str.startswith("https://"):
                req = urllib.request.Request(logo_str, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=5) as response:
                    image_stream = io.BytesIO(response.read())
            elif os.path.exists(logo_str):
                image_stream = logo_str

            if image_stream:
                logo_left = Inches(slide_width_in - 1.5) if not is_cover else Inches((slide_width_in / 2) - 0.6)
                logo_top = Inches(0.2) if not is_cover else Inches(0.8)
                logo_width = Inches(1.2) if not is_cover else Inches(1.4)
                slide.shapes.add_picture(image_stream, logo_left, logo_top, width=logo_width)
        except Exception as err:
            logger.warning("Failed to render custom brand logo: %s", err)


# ---------------------------------------------------------------------
# PptRenderer
# ---------------------------------------------------------------------

class PptRenderer:
    def __init__(self, template_file: str = DEFAULT_TEMPLATE_FILE):
        self.template_file = template_file

    def auto_select_layout(self, slide_spec: SlideSpec, slide_index: int) -> str:
        plugin_types = {p.type for p in slide_spec.plugins}
        content_types = plugin_types - {"notes"}

        if slide_spec.layout:
            return slide_spec.layout

        if slide_index == 0 and content_types == {"text"}:
            return "title_slide"

        mixed_kinds = {"paragraph", "bullets", "image", "chart", "table"}
        if len(content_types & mixed_kinds) >= 2:
            return "mixed_content_slide"

        if "chart" in content_types:
            return "chart_slide"
        if "image" in content_types:
            return "image_slide"
        if "table" in content_types:
            return "table_slide"
        if content_types == {"bullets"}:
            return "bullets_slide"
        if "paragraph" in content_types:
            return "title_content"
        if "text" in content_types:
            return "title_content"

        return "title_content"

    def render(
        self,
        plan: PresentationPlan,
        content_theme: Optional[str] = None,
        visual_style: Optional[str] = None,
    ) -> Presentation:
        prs = ensure_template_prs(self.template_file)
        layout_registry = get_layout_registry(self.template_file)
        active_theme = plan.theme or content_theme

        for idx, slide_spec in enumerate(plan.slides):
            layout_key = self.auto_select_layout(slide_spec, idx)
            layout_spec = layout_registry.get(layout_key, LayoutSpec(layout_index=0))
            layout_index = layout_spec.layout_index
            if layout_index >= len(prs.slide_layouts):
                layout_index = 0

            slide_layout = prs.slide_layouts[layout_index]
            slide = prs.slides.add_slide(slide_layout)

            for sp in list(slide.placeholders):
                try:
                    sp._element.getparent().remove(sp._element)
                except Exception:
                    pass

            apply_background_theme(slide, active_theme, visual_style=visual_style)

            palette = get_theme_palette(active_theme)
            if plan.use_custom_brand or plan.brand_color or plan.brand_secondary_color:
                if plan.brand_color:
                    palette["accent"] = hex_to_rgb(plan.brand_color)
                if plan.brand_secondary_color:
                    palette["background"] = hex_to_rgb(plan.brand_secondary_color)

            is_cover = idx == 0 or slide_spec.layout in {"title_slide", "section_slide"}
            current_y = 2.0 if is_cover else 0.35
            slide_width_in = float(prs.slide_width / Inches(1))
            left_margin = 0.6
            content_width = max(6.0, slide_width_in - (left_margin * 2.0))

            add_brand_elements_to_slide(slide, plan, slide_width_in, is_cover)

            if visual_style in {"corporate", "modern_gradient"} and not is_cover:
                try:
                    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(0.0), Inches(slide_width_in), Inches(0.06))
                    top_bar.fill.solid()
                    top_bar.fill.fore_color.rgb = palette["accent"]
                    top_bar.line.fill.background()
                except Exception:
                    pass

            bg_is_light = is_light_color(palette["background"])
            badge_color = palette.get("badge") or palette["accent"]
            if bg_is_light and is_light_color(badge_color):
                badge_color = palette["text"]

            if not is_cover:
                footer_box = slide.shapes.add_textbox(Inches(slide_width_in - 2.5), Inches(7.0), Inches(2.0), Inches(0.3))
                p_b = footer_box.text_frame.paragraphs[0]
                p_b.text = f"SLIDE {idx + 1} OF {len(plan.slides)}"
                p_b.font.size = Pt(8)
                p_b.font.bold = True
                p_b.alignment = PP_ALIGN.RIGHT
                p_b.font.color.rgb = badge_color

            title_text = slide_spec.title or (plan.title if idx == 0 else "")
            raw_t_align = str(slide_spec.title_align or "auto").lower().strip()
            raw_t_valign = str(slide_spec.title_valign or "auto").lower().strip()
            raw_s_align = str(slide_spec.subtitle_align or "auto").lower().strip()
            raw_s_valign = str(slide_spec.subtitle_valign or "auto").lower().strip()

            align_map = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY, "left": PP_ALIGN.LEFT}
            v_align_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

            if is_cover:
                if raw_t_valign in {"bottom", "down"}:
                    current_y = 5.3 if slide_spec.subtitle else 5.8
                elif raw_t_valign in {"top"}:
                    current_y = 0.8
                elif raw_t_valign in {"middle", "center"}:
                    current_y = 2.4
            else:
                if raw_t_valign in {"bottom", "down"}:
                    current_y = 5.8 if slide_spec.subtitle else 6.2
                elif raw_t_valign in {"middle", "center"}:
                    current_y = 3.2

            if title_text:
                title_color = hex_to_rgb(slide_spec.title_color) if slide_spec.title_color else palette["text"]
                title_bold = slide_spec.title_bold if slide_spec.title_bold is not None else True

                if is_cover:
                    # Dynamic title sizing for title cover to prevent line wrap overlap with subtitle
                    if len(title_text) > 45:
                        default_title_size = 36
                    elif len(title_text) > 28:
                        default_title_size = 40
                    else:
                        default_title_size = 46
                    title_font_size = slide_spec.title_font_size or default_title_size

                    auto_h = "center"
                    if raw_t_align in {"auto", "none", ""}:
                        raw_t_align = auto_h
                    t_align = align_map.get(raw_t_align, PP_ALIGN.CENTER)

                    # Estimate wrapped line count (at ~12 in width, ~28 chars per line at 40-46pt)
                    est_lines = max(1, math.ceil(len(title_text) / 28.0))
                    box_h = max(0.95, round(est_lines * (title_font_size / 72.0 * 1.35), 2))

                    t_box = slide.shapes.add_textbox(Inches(left_margin), Inches(current_y), Inches(content_width), Inches(box_h))
                    tf_t = t_box.text_frame
                    tf_t.word_wrap = True
                    if raw_t_valign in v_align_map:
                        tf_t.vertical_anchor = v_align_map[raw_t_valign]
                    p_t = tf_t.paragraphs[0]
                    p_t.text = title_text
                    p_t.alignment = t_align
                    set_run_style(p_t.runs[0] if p_t.runs else p_t.add_run(), font_size=title_font_size, bold=title_bold, color=title_color)
                    
                    # Generous spacing after title to position subtitle safely below all wrapped title lines
                    current_y += (box_h + 0.25)
                else:
                    default_title_size = 25 if len(title_text) > 45 else 29
                    title_font_size = slide_spec.title_font_size or default_title_size

                    auto_h = "left"
                    if raw_t_align in {"auto", "none", ""}:
                        raw_t_align = auto_h
                    t_align = align_map.get(raw_t_align, PP_ALIGN.LEFT)

                    est_lines = max(1, math.ceil(len(title_text) / 45.0))
                    box_h = max(0.65, round(est_lines * (title_font_size / 72.0 * 1.25), 2))

                    t_box = slide.shapes.add_textbox(Inches(left_margin), Inches(current_y), Inches(content_width), Inches(box_h))
                    tf_t = t_box.text_frame
                    tf_t.word_wrap = True
                    if raw_t_valign in v_align_map:
                        tf_t.vertical_anchor = v_align_map[raw_t_valign]
                    p_t = tf_t.paragraphs[0]
                    p_t.text = title_text
                    p_t.alignment = t_align
                    set_run_style(p_t.runs[0] if p_t.runs else p_t.add_run(), font_size=title_font_size, bold=title_bold, color=title_color)
                    current_y += (box_h + 0.10)

            # Subtitle Rendering (with duplicate title suppression)
            subtitle_text = (slide_spec.subtitle or "").strip()
            if title_text and subtitle_text and subtitle_text.lower() == title_text.lower():
                subtitle_text = ""

            if subtitle_text:
                sub_font_size = slide_spec.subtitle_font_size or (20 if is_cover else 18)
                sub_color = hex_to_rgb(slide_spec.subtitle_color) if slide_spec.subtitle_color else RGBColor(148, 163, 184)

                auto_s_h = "center" if is_cover else "left"
                if raw_s_align in {"auto", "none", ""}:
                    raw_s_align = auto_s_h

                s_align = align_map.get(raw_s_align, PP_ALIGN.CENTER if is_cover else PP_ALIGN.LEFT)

                # Estimate subtitle wrapped lines
                est_sub_lines = max(1, math.ceil(len(subtitle_text) / 50.0))
                sub_box_h = max(0.45, round(est_sub_lines * (sub_font_size / 72.0 * 1.25), 2))

                sub_box = slide.shapes.add_textbox(Inches(left_margin), Inches(current_y), Inches(content_width), Inches(sub_box_h))
                tf_s = sub_box.text_frame
                tf_s.word_wrap = True
                if raw_s_valign in v_align_map:
                    tf_s.vertical_anchor = v_align_map[raw_s_valign]
                p_s = tf_s.paragraphs[0]
                p_s.text = subtitle_text
                p_s.alignment = s_align
                set_run_style(p_s.runs[0] if p_s.runs else p_s.add_run(), font_size=sub_font_size, bold=False, color=sub_color)
                current_y += (sub_box_h + 0.15)

            current_y += 0.05

            for plugin in slide_spec.plugins:
                handler = PLUGIN_REGISTRY.get(plugin.type)
                if handler is None:
                    continue
                plugin_data = plugin.data.copy() if isinstance(plugin.data, dict) else {}
                if "slide_title" not in plugin_data and slide_spec.title:
                    plugin_data["slide_title"] = slide_spec.title
                plugin_data["theme_name"] = active_theme
                if (slide_spec.layout == "mixed_content_slide" or len(slide_spec.plugins) >= 2) and "box" in plugin_data:
                    handler.apply(slide, plugin_data, theme_name=active_theme)
                else:
                    try:
                        next_y = handler.apply_with_y(slide, plugin_data, current_y=current_y, left_margin=left_margin, content_width=content_width, palette=palette, theme_name=active_theme)
                    except TypeError:
                        next_y = handler.apply_with_y(slide, plugin_data, current_y=current_y, left_margin=left_margin, content_width=content_width, palette=palette)
                    if next_y is not None and next_y > current_y:
                        current_y = next_y

        return prs
